"""
embed_images.py
Embeds the 7 dashboard PNG diagrams into Section 8 of reports.docx.
Run: python embed_images.py
"""

import os
import copy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml   import OxmlElement

BASE    = os.path.dirname(__file__)
DOCPATH_IN  = os.path.join(BASE, "reports.docx")
DOCPATH_OUT = os.path.join(BASE, "reports_with_diagrams.docx")
IMGDIR  = os.path.join(BASE, "dashboard_images")

IMAGES = [
    ("01_kpi_dashboard.png",
     "Figure 1 — Page 1: KPI Dashboard  |  Five KPI cards + donut chart, subscription bar, region bar, fee histogram"),
    ("02_eda_demographics.png",
     "Figure 2 — EDA Demographics Tab  |  Age histogram, gender bar, device donut, genre bar — all by churn"),
    ("03_eda_behaviour_financials.png",
     "Figure 3 — EDA Viewing Behaviour & Financials  |  Box plots, scatter, payment method churn bars"),
    ("04_correlation_heatmap.png",
     "Figure 4 — Feature Correlation Heatmap  |  Pearson correlations across all 13 features"),
    ("05_predict_churn_page.png",
     "Figure 5 — Predict Churn Page  |  Gauge chart (73%), HIGH RISK badge, risk signals, thresholds"),
    ("06_model_performance.png",
     "Figure 6 — Model Performance  |  Accuracy comparison bars + feature importance chart (GradientBoost)"),
    ("07_architecture_flow.png",
     "Figure 7 — System Architecture  |  End-to-end data flow: CSV → train.py → models → FastAPI → Streamlit"),
]

doc = Document(DOCPATH_IN)
body = doc.element.body

# ── Find index of "8. Dashboard Overview" paragraph ───────────────────────────
anchor_idx = None
for i, para in enumerate(doc.paragraphs):
    if "8. Dashboard Overview" in para.text:
        anchor_idx = i
        break

if anchor_idx is None:
    raise RuntimeError("Could not find '8. Dashboard Overview' in the document.")

print(f"Anchor at paragraph index {anchor_idx}: '{doc.paragraphs[anchor_idx].text}'")

# ── We'll append temp paragraphs to the doc end, then move them after anchor ──
# This ensures the paragraph objects have full .part context for images.

def make_spacer(doc):
    p = doc.add_paragraph("")
    return p

def make_image_para(doc, img_path, width_inches=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_inches))
    return p

def make_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    return p

# Build groups of (spacer, image, caption, spacer) as XML elements
# then splice them in right after the anchor paragraph element.

anchor_elem = doc.paragraphs[anchor_idx]._p

# Insert in forward order: collect all new XML elements, then insert them
new_elements = []

for img_file, caption in IMAGES:
    img_path = os.path.join(IMGDIR, img_file)
    if not os.path.exists(img_path):
        print(f"  SKIP (not found): {img_file}")
        continue

    sp1  = make_spacer(doc)
    imgp = make_image_para(doc, img_path)
    cap  = make_caption(doc, caption)
    sp2  = make_spacer(doc)

    new_elements += [sp1._p, imgp._p, cap._p, sp2._p]
    print(f"  Prepared: {img_file}")

# Remove these elements from the end of body (they were appended there)
# and re-insert them after the anchor element
for elem in new_elements:
    body.remove(elem)

# Insert after anchor in correct order
insert_after = anchor_elem
for elem in new_elements:
    insert_after.addnext(elem)
    insert_after = elem

doc.save(DOCPATH_OUT)
print(f"\nSaved: {DOCPATH_OUT}")
print(f"Total paragraphs: {len(doc.paragraphs)}")
